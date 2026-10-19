from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from config import CELL_MARKER, ROW_MARKER


def _cell_id(sheet: str, row: int, column: int) -> str:
    return f"{sheet}!{get_column_letter(column)}{row}"


def read_xlsx(path: str | Path) -> list[dict[str, Any]]:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx files are supported.")
    workbook = load_workbook(path, data_only=True)
    cells: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        merge_lookup: dict[tuple[int, int], str] = {}
        for merged in sheet.merged_cells.ranges:
            rng = str(merged)
            min_row, min_col, max_row, max_col = (
                merged.min_row,
                merged.min_col,
                merged.max_row,
                merged.max_col,
            )
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    merge_lookup[(r, c)] = rng
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                text = str(value).strip()
                if not text:
                    continue
                r, c = cell.row, cell.column
                cells.append(
                    {
                        "cell_id": _cell_id(sheet.title, r, c),
                        "sheet": sheet.title,
                        "row": r,
                        "column": c,
                        "original_text": text,
                        "english_text": None,
                        "merge_range": merge_lookup.get((r, c)),
                    }
                )
    return sort_cells(cells)


def sort_cells(cells: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(cells, key=lambda x: (x["sheet"], x["row"], x["column"]))


def build_model_input(
    cells: list[dict[str, Any]], use_english: bool = True
) -> tuple[str, list[dict[str, Any]]]:
    ordered = sort_cells(cells)
    parts: list[str] = []
    spans: list[dict[str, Any]] = []
    current_sheet = None
    current_row = None
    cursor = 0

    def append_token(token: str) -> None:
        nonlocal cursor
        if parts:
            parts.append(" ")
            cursor += 1
        parts.append(token)
        cursor += len(token)

    for cell in ordered:
        if cell["sheet"] != current_sheet or cell["row"] != current_row:
            append_token(ROW_MARKER)
            current_sheet = cell["sheet"]
            current_row = cell["row"]
        append_token(CELL_MARKER)
        text = cell.get("english_text") if use_english else None
        if not text:
            text = cell["original_text"]
        start = cursor
        if parts:
            parts.append(" ")
            cursor += 1
            start = cursor
        parts.append(text)
        cursor += len(text)
        spans.append(
            {
                "start": start,
                "end": cursor,
                "cell_id": cell["cell_id"],
                "sheet": cell["sheet"],
                "row": cell["row"],
                "column": cell["column"],
                "text": text,
            }
        )
    return "".join(parts), spans
